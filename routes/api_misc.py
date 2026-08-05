from flask import Blueprint, request, jsonify, send_file
import pandas as pd
import json
import io
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from database import get_db_connection

api_misc = Blueprint('api_misc', __name__)

@api_misc.route('/api/list-mesin', methods=['GET'])
def get_list_mesin():
    try:
        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT DISTINCT mesin FROM qc_records WHERE mesin IS NOT NULL ORDER BY mesin ASC")
        mesin_list = [row['mesin'] for row in cursor.fetchall()]
        cursor.close()
        conn.close()
        if not mesin_list:
            mesin_list = ['14']
        return jsonify({"status": "success", "data": mesin_list}), 200
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 400

@api_misc.route('/api/export-excel', methods=['GET'])
def export_excel():
    start_date = request.args.get('start_date')
    end_date = request.args.get('end_date')
    id_produk = request.args.get('id_produk')

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    if start_date and end_date:
        query = """
            SELECT q.*, p.nama_produk, p.parameter_dinamis as standar_parameter 
            FROM qc_records q
            LEFT JOIN master_produk p ON q.id_produk = p.id
            WHERE q.id_produk = %s AND q.tanggal BETWEEN %s AND %s 
            ORDER BY q.tanggal ASC, q.id ASC
        """
        cursor.execute(query, (id_produk, start_date, end_date))
    else:
        query = """
            SELECT q.*, p.nama_produk, p.parameter_dinamis as standar_parameter 
            FROM qc_records q
            LEFT JOIN master_produk p ON q.id_produk = p.id
            WHERE q.id_produk = %s
            ORDER BY q.tanggal ASC, q.id ASC
        """
        cursor.execute(query, (id_produk,))
    data = cursor.fetchall()
    
    if not data:
        cursor.close()
        conn.close()
        return "TIDAK ADA DATA DI RENTANG TANGGAL TERSEBUT.", 404

    df = pd.DataFrame(data)
    
    import re
    def format_shift(row):
        s = str(row['shift']) if pd.notnull(row['shift']) else ''
        g = str(row['grup']) if pd.notnull(row['grup']) else ''
        if re.search('[a-zA-Z]', s):
            return s
        return s + g
        
    df['Shift_Gabung'] = df.apply(format_shift, axis=1)

    dynamic_cols = []
    status_cols = []
    if 'parameter_dinamis' in df.columns:
        parsed_params = df['parameter_dinamis'].apply(lambda val: json.loads(val) if pd.notnull(val) and val != '' else {})
        df_params = pd.json_normalize(parsed_params)
        df = pd.concat([df.drop('parameter_dinamis', axis=1), df_params], axis=1)
        dynamic_cols = list(df_params.columns)
        
        for col in dynamic_cols:
            status_col_name = f"Status {col} Std"
            status_cols.append(status_col_name)
            
            def get_status(row, c=col):
                try:
                    val = float(row[c])
                    std_json = json.loads(row['standar_parameter'])
                except:
                    return "-"
                    
                for sp in std_json:
                    if sp.get('name') == c:
                        lsl = sp.get('lcl')
                        usl = sp.get('ucl')
                        is_ok = True
                        if lsl is not None and str(lsl).strip() != "":
                            if val < float(lsl): is_ok = False
                        if usl is not None and str(usl).strip() != "":
                            if val > float(usl): is_ok = False
                        return "OK" if is_ok else "OUT"
                return "-"
                
            df[status_col_name] = df.apply(get_status, axis=1)

    summary_grouped = {}
    for nama_prod in df['nama_produk'].dropna().unique():
        df_prod = df[df['nama_produk'] == nama_prod]
        
        if 'standar_parameter' in df_prod.columns:
            standar_str = df_prod['standar_parameter'].iloc[-1] if not df_prod.empty and not pd.isna(df_prod['standar_parameter'].iloc[-1]) else '[]'
            try:
                standar_json = json.loads(standar_str)
            except:
                standar_json = []
                
            for std_param in standar_json:
                p_name = std_param.get('name')
                if p_name and p_name in df_prod.columns:
                    df_prod_param = pd.to_numeric(df_prod[p_name], errors='coerce')
                    p_mean = df_prod_param.mean()
                    p_std = df_prod_param.std()
                    if pd.isna(p_std): p_std = 0
                    
                    lsl_dyn = std_param.get('lcl')
                    if lsl_dyn is None: lsl_dyn = 0
                    usl_dyn = std_param.get('ucl')
                    if usl_dyn is None: usl_dyn = 0
                    
                    p_calc_lcl = p_mean - (3 * p_std)
                    p_calc_ucl = p_mean + (3 * p_std)
                    p_status = "Capable" if (p_calc_lcl >= lsl_dyn and p_calc_ucl <= usl_dyn) else "Out of Control"
                    
                    if nama_prod not in summary_grouped:
                        summary_grouped[nama_prod] = []
                    
                    summary_grouped[nama_prod].append({
                        "Parameter Name": p_name,
                        "Mean (Rata-rata)": round(p_mean, 3) if not pd.isna(p_mean) else "-",
                        "Std Dev (Sigma)": round(p_std, 3),
                        "LSL (Spec Limit Bawah)": lsl_dyn,
                        "USL (Spec Limit Atas)": usl_dyn,
                        "Calculated LCL (Control Limit Bawah)": round(p_calc_lcl, 3) if not pd.isna(p_calc_lcl) else "-",
                        "Calculated UCL (Control Limit Atas)": round(p_calc_ucl, 3) if not pd.isna(p_calc_ucl) else "-",
                        "Status": p_status
                    })

    if 'standar_parameter' in df.columns:
        df = df.drop('standar_parameter', axis=1)
    
    df['Kesimpulan Proses'] = 'Proses Stabil'

    kolom_baru = {'tanggal': 'Tanggal', 'nama_produk': 'Produk', 'Shift_Gabung': 'Shift', 'mesin': 'Mesin', 'waktu': 'Waktu'}
    df_renamed = df.rename(columns=kolom_baru)
    urutan_kolom = list(kolom_baru.values()) + dynamic_cols + ['Kesimpulan Proses'] + status_cols
    kolom_tersedia = [col for col in urutan_kolom if col in df_renamed.columns]
    df_final = df_renamed[kolom_tersedia]
    
    df_final.insert(0, 'No', range(1, len(df_final) + 1))
    df_final['Tanggal'] = pd.to_datetime(df_final['Tanggal']).dt.strftime('%d %b %Y')

    output_file = io.BytesIO()
    
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        def styling_sheet(worksheet, header_row=2):
            font_header = Font(bold=True, color="FFFFFF")
            align_center = Alignment(horizontal="center", vertical="center")
            border_thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
            warna_biru = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            warna_merah = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
            warna_hijau = PatternFill(start_color="548235", end_color="548235", fill_type="solid")

            for col_num, column_cells in enumerate(worksheet.columns, 1):
                nama_kolom = column_cells[header_row - 1].value 
                panjang_maksimal = max((len(str(cell.value)) for cell in column_cells if cell.value is not None and cell.row >= header_row), default=10)
                worksheet.column_dimensions[column_cells[0].column_letter].width = min(panjang_maksimal + 2, 50)
                
                for cell in column_cells:
                    if cell.row == header_row:
                        cell.font = font_header
                        cell.alignment = align_center
                        cell.border = border_thin
                        if nama_kolom in ['Status Tebal', 'Kesimpulan Proses']: cell.fill = warna_merah
                        elif nama_kolom and 'Std' in str(nama_kolom): cell.fill = warna_hijau
                        else: cell.fill = warna_biru
                    elif cell.row > header_row:
                        cell.border = border_thin
                        cell.alignment = align_center

        daftar_produk_unik = df_final['Produk'].dropna().unique()
        for idx, nama_prod in enumerate(daftar_produk_unik):
            df_filter = df_final[df_final['Produk'] == nama_prod].dropna(axis=1, how='all')
            sheet_name = f"DATA_{nama_prod}"[:31]
            import re
            sheet_name = re.sub(r'[\\/*?:\[\]]', '', sheet_name)
            if not sheet_name: sheet_name = f"Produk_{idx}"
            df_filter.to_excel(writer, index=False, sheet_name=sheet_name, startrow=1)
            ws_data = writer.sheets[sheet_name]
            ws_data.cell(row=1, column=1, value=f"MONITORING {str(nama_prod).upper()}").font = Font(bold=True)
            styling_sheet(ws_data, header_row=2)

        if summary_grouped:
            ws_std = writer.book.create_sheet('Standar Parameter')
            font_header = Font(bold=True, color="FFFFFF")
            warna_biru = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            align_center = Alignment(horizontal="center", vertical="center")
            border_thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
            
            ws_std.cell(row=1, column=1, value="Parameter").font = font_header
            ws_std.cell(row=1, column=1).fill = warna_biru
            ws_std.cell(row=1, column=1).alignment = align_center
            ws_std.cell(row=1, column=1).border = border_thin
            
            ws_std.cell(row=1, column=2, value="LSL").font = font_header
            ws_std.cell(row=1, column=2).fill = warna_biru
            ws_std.cell(row=1, column=2).alignment = align_center
            ws_std.cell(row=1, column=2).border = border_thin
            
            ws_std.cell(row=1, column=3, value="USL").font = font_header
            ws_std.cell(row=1, column=3).fill = warna_biru
            ws_std.cell(row=1, column=3).alignment = align_center
            ws_std.cell(row=1, column=3).border = border_thin
            
            row_idx = 2
            for prod, params in summary_grouped.items():
                for param in params:
                    c1 = ws_std.cell(row=row_idx, column=1, value=param["Parameter Name"])
                    c2 = ws_std.cell(row=row_idx, column=2, value=param["LSL (Spec Limit Bawah)"])
                    c3 = ws_std.cell(row=row_idx, column=3, value=param["USL (Spec Limit Atas)"])
                    for c in [c1, c2, c3]:
                        c.border = border_thin
                        c.alignment = align_center
                    row_idx += 1

        if summary_grouped:
            ws = writer.book.create_sheet('Summary')
            font_header = Font(bold=True, color="FFFFFF")
            warna_hijau_tua = PatternFill(start_color="375623", end_color="375623", fill_type="solid")
            align_center = Alignment(horizontal="center", vertical="center")
            align_left = Alignment(horizontal="left", vertical="center")
            border_thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
            row_idx = 1
            headers = ["Variable", "CL", "Std Dev", "UCL", "LCL"]
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=header)
                cell.font = font_header
                cell.fill = warna_hijau_tua
                cell.alignment = align_center
                cell.border = border_thin
            row_idx += 1
            for prod, params in summary_grouped.items():
                for param in params:
                    row_data = [
                        param["Parameter Name"], param["Mean (Rata-rata)"], param["Std Dev (Sigma)"],
                        param["Calculated UCL (Control Limit Atas)"], param["Calculated LCL (Control Limit Bawah)"]
                    ]
                    for col_idx, val in enumerate(row_data, 1):
                        cell = ws.cell(row=row_idx, column=col_idx, value=val)
                        cell.border = border_thin
                        cell.alignment = align_left if col_idx == 1 else align_center
                    row_idx += 1

    cursor.close()
    conn.close()

    output_file.seek(0)
    nama_file = f"Laporan_QC_{start_date}_sd_{end_date}.xlsx" if start_date and end_date else "Laporan_QC_All_Time.xlsx"

    return send_file(
        output_file,
        as_attachment=True,
        download_name=nama_file,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
